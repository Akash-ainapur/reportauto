from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
import os
import json
import re

# -------------------------------------------------
# Add page border
# -------------------------------------------------
def add_page_border(section):
    sectPr = section._sectPr
    for child in list(sectPr):
        if child.tag == qn("w:pgBorders"):
            sectPr.remove(child)

    pgBorders = OxmlElement("w:pgBorders")
    pgBorders.set(qn("w:offsetFrom"), "page")

    for side in ["top", "left", "bottom", "right"]:
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:space"), "24")
        border.set(qn("w:color"), "000000")
        pgBorders.append(border)

    sectPr.append(pgBorders)


# -------------------------------------------------
# ✅ FIXED replacement (run-safe & bold-safe)
# -------------------------------------------------
def replace_chunk_with_format(paragraph, old_text, new_text, is_regex=False, 
                              bold=False, size=None, color=None):
    if not paragraph.runs:
        return

    # 1. Capture state of existing runs
    run_data = []
    for r in paragraph.runs:
        run_data.append({
            "text": r.text,
            "bold": r.bold,
            "italic": r.italic,
            "underline": r.underline,
            "font_name": r.font.name,
            "font_size": r.font.size,
            "font_color": r.font.color.rgb,
            "original_run": r
        })
    
    full_text = "".join(d["text"] for d in run_data)
    
    # 2. Find match
    match_start = -1
    match_end = -1
    
    if is_regex:
        match = re.search(old_text, full_text, flags=re.IGNORECASE)
        if match:
            match_start = match.start()
            match_end = match.end()
    else:
        match_start = full_text.find(old_text)
        if match_start != -1:
            match_end = match_start + len(old_text)
            
    if match_start == -1:
        return

    # 3. Clear existing runs
    for r in paragraph.runs:
        r.text = ""
        
    # 4. Rebuild runs
    current_pos = 0
    
    for r_item in run_data:
        r_text = r_item["text"]
        r_len = len(r_text)
        r_end = current_pos + r_len
        
        # Helper to create run with copied props
        def add_run_copy_props(text):
            nr = paragraph.add_run(text)
            nr.bold = r_item["bold"]
            nr.italic = r_item["italic"]
            nr.underline = r_item["underline"]
            if r_item["font_name"]: nr.font.name = r_item["font_name"]
            if r_item["font_size"]: nr.font.size = r_item["font_size"]
            if r_item["font_color"]: nr.font.color.rgb = r_item["font_color"]
            return nr

        # Case A: Run is completely before match
        if r_end <= match_start:
             add_run_copy_props(r_text)
             
        # Case B: Run is completely after match
        elif current_pos >= match_end:
             add_run_copy_props(r_text)
             
        # Case C: Overlap
        else:
            # Sub-case: Part before match
            if current_pos < match_start:
                start_offset = 0
                end_offset = match_start - current_pos
                text_chunk = r_text[start_offset:end_offset]
                if text_chunk:
                    add_run_copy_props(text_chunk)
            
            # Sub-case: The Match Itself (Insert ONLY ONCE when we hit the start)
            # We insert the REPLACEMENT at match_start.
            if current_pos <= match_start < r_end:
                 target_run = paragraph.add_run(new_text)
                 # Apply requested formatting or inherit
                 target_run.bold = True if bold else r_item["bold"]
                 if size: target_run.font.size = Pt(size)
                 if color: target_run.font.color.rgb = color
                 # Inherit others if not specified? 
                 # For now, bold is explicit override.
            
            # Sub-case: Part after match
            if r_end > match_end:
                start_offset = match_end - current_pos
                if start_offset < 0: start_offset = 0
                text_chunk = r_text[start_offset:]
                if text_chunk:
                    add_run_copy_props(text_chunk)
                    
        current_pos += r_len


# -------------------------------------------------
# Main Generator
# -------------------------------------------------
def update_acknowledgement_page(input_path):

    current_script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(os.path.dirname(os.path.dirname(current_script_dir)))

    output_dir = os.path.join(project_root, "output")
    os.makedirs(output_dir, exist_ok=True)

    filename = os.path.basename(input_path)
    base, ext = os.path.splitext(filename)
    output_path = os.path.join(output_dir, f"{base}_filled_new{ext}")

    doc = Document(input_path)

    # Load JSON
    json_path = os.path.join(project_root, "backend", "data", "full_report_template.json")
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    approvers = data["people"]["approvers"]
    guide = data["people"]["guide"][0]

    # -------------------------------------------------
    # ✅ Replacement configuration (UNCHANGED LOGIC)
    # -------------------------------------------------
    replacements = [
        (r"Dr\.\s*Ajayan\s*K\s*R", approvers["principal"]["name"], True, True, None, None),
        (r"Principal\s+of\s+MVJCE", "Principal of MVJCE", True, True, None, None),

        (r"Dr\.\s*I\s*Hameem\s*Shanavas", approvers["dean"]["name"], True, True, None, None),
        (r"Dean,\s*School\s+of\s+CSE,\s*MVJCE", "Dean, School of CSE, MVJCE", True, True, None, None),

        (r"Mr\.\s*Kumar\s*R", approvers["COE"]["name"], True, True, None, None),

        (r"Controller\s+Of\s+Examinations", "Controller Of Examinations", True, True, None, None),

        (r"Dr\.\s*Kiran\s*Babu\s*T\s*S", approvers["hod"]["name"], True, True, None, None),

        (r"Mrs\.\s*Sumanjali\s*M", guide["name"], True, True, None, None),
        (r"Assistant\s+Professor|HEAD\s+WIZARD", guide["designation"], True, True, None, None),

        (r"CSE\s+Department|ECE\s+Department", guide["department"], True, True, None, None),
    ]


    def process_paragraphs(paragraphs):
        for p in paragraphs:
            for search, replace, is_re, bold, size, color in replacements:
                replace_chunk_with_format(p, search, replace, is_re, bold, size, color)

    process_paragraphs(doc.paragraphs)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                process_paragraphs(cell.paragraphs)

    add_page_border(doc.sections[0])
    doc.save(output_path)

    print(f"✅ Acknowledgement generated correctly: {output_path}")


# -------------------------------------------------
# Run
# -------------------------------------------------
if __name__ == "__main__":
    current_script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(os.path.dirname(os.path.dirname(current_script_dir)))
    input_file = os.path.join(
        project_root,
        "backend",
        "templates",
        "template_indivisual",
        "Aknowledgement.docx"
    )

    if os.path.exists(input_file):
        update_acknowledgement_page(input_file)
    else:
        print(f"Template file not found at {input_file}")
