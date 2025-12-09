from docx import Document
import os
import sys

# Add project root to path to run the generator if needed, but we can just run it via command line
# or import the function. Let's run the main function.

# Adjust path to import generator script
current_dir = os.path.dirname(os.path.abspath(__file__))
project_root = os.path.dirname(os.path.dirname(os.path.dirname(current_dir)))
sys.path.append(os.path.join(project_root, "backend", "scripts", "generator"))

from page_01_front_page import update_project_doc_with_dummies

def verify():
    # 1. Run the generator
    input_file = os.path.join(project_root, "backend", "templates", "template_indivisual", "front_page.docx")
    print(f"Running generator on: {input_file}")
    update_project_doc_with_dummies(input_file)
    
    # 2. Check the output
    output_path = os.path.join(project_root, "output", "front_page_filled_verification.docx")
    if not os.path.exists(output_path):
        print("FAIL: Output file not created.")
        return

    doc = Document(output_path)
    
    # Check Designation Casing
    # Expected: "Head Wizard, Dept Of Magic"
    # Original Input Data: designation="HEAD WIZARD", department="DEPT OF MAGIC"
    expected_designation = "Head Wizard, Dept Of Magic"
    items_found = 0
    
    for para in doc.paragraphs:
        if expected_designation in para.text:
            print(f"PASS: Found Title Case designation: '{expected_designation}'")
            items_found += 1
            break
    
    if items_found == 0:
        print(f"FAIL: Did not find designation '{expected_designation}'")
        # Print close matches?
        for para in doc.paragraphs:
            if "Wizard" in para.text:
                print(f"  Found similar: '{para.text}'")

    # Check Guide Name Bolding
    # Expected Name: ">>> PROF. DUMBLEDORE <<<"
    expected_name = ">>> PROF. DUMBLEDORE <<<"
    name_found = False
    
    for para in doc.paragraphs:
        if expected_name in para.text:
            # Check runs for this paragraph
            # We expect the name part to be in a run with bold=True
            for run in para.runs:
                if expected_name in run.text:
                    if run.bold:
                        print(f"PASS: Guide Name '{run.text}' is BOLD.")
                        name_found = True
                    else:
                        print(f"FAIL: Guide Name '{run.text}' is present but NOT BOLD.")
            
            if not name_found:
                 print(f"FAIL: Guide Name found but could not verify bold status (maybe split across runs improperly?)")
            break

    if not name_found:
        print("FAIL: Did not find Guide Name paragraph.")

if __name__ == "__main__":
    verify()
