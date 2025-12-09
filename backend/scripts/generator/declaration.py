from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import json

# -------------------------------------------------
# Add page border
# -------------------------------------------------
def add_page_border(section):
    sectPr = section._sectPr
    for child in list(sectPr):
        if child.tag == qn("w:pgBorders"):
            sectPr.remove(child)
    pgBorders = OxmlElement("w:pgBorders")
    pgBorders.set(qn("w:offsetFrom"), "page")
    for side in ["top", "left", "bottom", "right"]:
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:space"), "24")
        border.set(qn("w:color"), "000000")
        pgBorders.append(border)
    sectPr.append(pgBorders)

# -------------------------------------------------
# Main Generator
# -------------------------------------------------
def update_declaration_page(input_path):
    # Path setup
    current_script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(os.path.dirname(os.path.dirname(current_script_dir)))
    
    # Establish output directory
    output_dir = os.path.join(project_root, "output")
    os.makedirs(output_dir, exist_ok=True)

    filename = os.path.basename(input_path)
    base, ext = os.path.splitext(filename)
    output_filename = f"{base}_filled_v3{ext}"
    output_path = os.path.join(output_dir, output_filename)

    doc = Document(input_path)

    # Load JSON
    json_path = os.path.join(project_root, "backend", "data", "full_report_template.json")
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    meta = data["global_metadata"]
    students = data["people"]["students"]

    from docx.shared import RGBColor
    import re

    def replace_chunk_with_format(paragraph, old_text, new_text, is_regex=False, 
                               bold=False, size=None, color=None):
        if not paragraph.runs:
            return

        # 1. Capture state of existing runs
        run_data = []
        for r in paragraph.runs:
            run_data.append({
                "text": r.text,
                "bold": r.bold,
                "italic": r.italic,
                "underline": r.underline,
                "font_name": r.font.name,
                "font_size": r.font.size,
                "font_color": r.font.color.rgb,
                "original_run": r
            })
        
        full_text = "".join(d["text"] for d in run_data)
        
        # 2. Find match
        match_start = -1
        match_end = -1
        
        if is_regex:
            match = re.search(old_text, full_text, flags=re.IGNORECASE)
            if match:
                match_start = match.start()
                match_end = match.end()
        else:
            match_start = full_text.find(old_text)
            if match_start != -1:
                match_end = match_start + len(old_text)
                
        if match_start == -1:
            return

        # 3. Clear existing runs
        for r in paragraph.runs:
            r.text = ""
            
        # 4. Rebuild runs
        current_pos = 0
        
        for r_item in run_data:
            r_text = r_item["text"]
            r_len = len(r_text)
            r_end = current_pos + r_len
            
            # Helper to create run with copied props
            def add_run_copy_props(text):
                nr = paragraph.add_run(text)
                nr.bold = r_item["bold"]
                nr.italic = r_item["italic"]
                nr.underline = r_item["underline"]
                if r_item["font_name"]: nr.font.name = r_item["font_name"]
                if r_item["font_size"]: nr.font.size = r_item["font_size"]
                if r_item["font_color"]: nr.font.color.rgb = r_item["font_color"]
                return nr

            # Case A: Run is completely before match
            if r_end <= match_start:
                 add_run_copy_props(r_text)
                 
            # Case B: Run is completely after match
            elif current_pos >= match_end:
                 add_run_copy_props(r_text)
                 
            # Case C: Overlap
            else:
                # Sub-case: Part before match
                if current_pos < match_start:
                    start_offset = 0
                    end_offset = match_start - current_pos
                    text_chunk = r_text[start_offset:end_offset]
                    if text_chunk:
                        add_run_copy_props(text_chunk)
                
                # Sub-case: The Match Itself (Insert ONLY ONCE when we hit the start)
                if current_pos <= match_start < r_end:
                     # Check if new_text is a LIST of segments for mixed formatting
                     if isinstance(new_text, list):
                         for seg in new_text:
                             seg_text = seg.get("text", "")
                             seg_bold = seg.get("bold", bold) # default to func arg
                             seg_color = seg.get("color", color) # default to func arg
                             
                             tr = paragraph.add_run(seg_text)
                             tr.bold = True if seg_bold else r_item["bold"]
                             if size: tr.font.size = Pt(size)
                             if seg_color: tr.font.color.rgb = seg_color
                     else:
                         # Standard single replacement
                         target_run = paragraph.add_run(new_text)
                         target_run.bold = True if bold else r_item["bold"]
                         if size: target_run.font.size = Pt(size)
                         if color: target_run.font.color.rgb = color
                
                # Sub-case: Part after match
                if r_end > match_end:
                    start_offset = match_end - current_pos
                    if start_offset < 0: start_offset = 0
                    text_chunk = r_text[start_offset:]
                    if text_chunk:
                        add_run_copy_props(text_chunk)
                        
            current_pos += r_len

    # Student Data Helpers (Safe Access)
    def get_stud_name(idx): return students[idx]["name"] if idx < len(students) else f"Student {idx+1}"
    def get_stud_usn(idx): return students[idx]["usn"] if idx < len(students) else f"USN {idx+1}"

    # Replacements Configuration
    # Format: (Search, Replace, Regex?, Bold?, Size?, Color?)
    replacements = [
        # College and Dept
        ("MVJ COLLEGE OF ENGINEERING", meta["college_name"], False, True, 18, None),
        ("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", meta["department_name"], False, True, 14, None),

        ("Computer science and Engineering", meta["department_short"], False, True, 12, None),
        
        # Semester
        (r"Sixth", meta["semester"], True, False, None, None),

        # Major -> Mini replacement
        (r"Major", meta["project_type"], True, False, None, None),

        # Student Names
        ("ARUNPRAKASH PATTAR", get_stud_name(0), False, False, None, None),
        ("HEMANTH S SAVANTH",  get_stud_name(1), False, False, None, None),
        ("CHINMAHESH C",       get_stud_name(2), False, False, None, None),
        ("BHAVYASHREE R",      get_stud_name(3), False, False, None, None),

        # Student USNs
        ("1MJ22CS217", get_stud_usn(0), False, False, None, None),
        ("1MJ23CS411", get_stud_usn(1), False, False, None, None),
        ("1MJ23CS406", get_stud_usn(2), False, False, None, None),
        ("1MJ22CS029", get_stud_usn(3), False, False, None, None),

        # Academic Year
        ("2024-2025", meta["academic_year"], False, False, None, None),
        ("2024 - 2025", meta["academic_year"], False, False, None, None), # Handle both formats
        
        # Dept Name in Body (using department_short as requested)
        ("Computer science and Engineering", meta.get("department_short", "Dept of Universal Logic"), False, False, None, None),

        # Project Title (Red, Robust Regex)
        # Replaced with Multi-Segment List for: Black Quotes + Red Title
        (r"[\"“]?\s*MVJ\s*NEXTGEN\s*EDUCLOUD[:\s-]+AI[:\s-]+BASED\s*ACADEMIC\s*RESOURCE\s*MANAGER\s*[\"”]?", 
         [
             {"text": "“", "bold": True, "color": None},       # Black Bold Quote
             {"text": meta["project_title"], "bold": True, "color": RGBColor(255, 0, 0)}, # Red Bold Title
             {"text": "”", "bold": True, "color": None},       # Black Bold Quote
         ], 
         True, True, None, RGBColor(255, 0, 0)),
    ]

    # Helper to run replacements on a list of paragraphs
    def process_paragraphs(paragraphs):
        for p in paragraphs:
            for search, replace, is_re, bold, size, color in replacements:
                replace_chunk_with_format(p, search, replace, is_re, bold, size, color)

    # 1. Process Body Paragraphs
    process_paragraphs(doc.paragraphs)
    
    # 2. Process Tables (Crucial for Student Names/USNs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                process_paragraphs(cell.paragraphs)

    # Add page border
    add_page_border(doc.sections[0])

    doc.save(output_path)
    print(f"✅ Declaration generated correctly: {output_path}")

# -------------------------------------------------
# Run
# -------------------------------------------------
if __name__ == "__main__":
    current_script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(os.path.dirname(os.path.dirname(current_script_dir)))
    input_file = os.path.join(project_root, "backend", "templates", "template_indivisual", "declaration.docx")
    
    if os.path.exists(input_file):
        update_declaration_page(input_file)
    else:
        print(f"Error: Template file not found at {input_file}")
