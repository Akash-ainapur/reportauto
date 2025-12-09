from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def generate_abstract_page(input_path=None):
    # If no input path is provided, creating a new blank document or using a template if it existed
    # Since we don't have a template for abstract yet, we'll create a simple one or load a blank one
    
    # Path setup
    current_script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(os.path.dirname(os.path.dirname(current_script_dir)))
    
    output_dir = os.path.join(project_root, "output")
    os.makedirs(output_dir, exist_ok=True)
    
    output_path = os.path.join(output_dir, "abstract_filled.docx")
    
    doc = Document()
    
    # Title
    p = doc.add_paragraph()
    run = p.add_run("ABSTRACT")
    run.bold = True
    run.font.size = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Placeholder content
    p = doc.add_paragraph()
    p.add_run("\n[Abstract Content Will Be Generated Here]\n")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Ensure it ends with a page break so the next document starts on a new page
    # doc.add_page_break() 
    # Actually, for docxcompose, it is safer to rely on the composer adding a break or section, 
    # but adding one here is a safety net. However, if composer adds one too, we get a blank page.
    # Let's fix the merging logic instead.

    
    doc.save(output_path)
    print(f"✅ Abstract generated: {output_path}")

if __name__ == "__main__":
    generate_abstract_page()
