from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os
import re

def replace_text_runs(paragraph, old, new):
    """Replace text even if split across runs."""
    full_text = "".join(run.text for run in paragraph.runs)
    if old in full_text:
        new_text = full_text.replace(old, new)
        # clear all runs and write new text to first run
        for run in paragraph.runs:
            run.text = ""
        paragraph.runs[0].text = new_text

def add_page_number(paragraph):
    """Adds dynamic PAGE field to footer paragraph."""
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')

    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

def update_project_doc_with_dummies(input_path):
    doc = Document(input_path)
    
    # Get the directory of the current script
    current_script_dir = os.path.dirname(os.path.abspath(__file__))
    # Go up 3 levels to get to project root (generator -> scripts -> backend -> root)
    project_root = os.path.dirname(os.path.dirname(os.path.dirname(current_script_dir)))
    
    output_dir = os.path.join(project_root, "output")
    os.makedirs(output_dir, exist_ok=True)
    
    filename = os.path.basename(input_path)
    base, ext = os.path.splitext(filename)
    output_filename = f"{base}_filled{ext}"
    output_path = os.path.join(output_dir, output_filename)

    # --- LOAD JSON DATA ---
    import json
    json_path = os.path.join(project_root, "backend", "data", "full_report_template.json")
    
    try:
        with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except FileNotFoundError:
        print(f"Error: JSON file not found at {json_path}")
        return

    meta = data["global_metadata"]
    guides = data["people"]["guides"][0]
    
    # --- TEXT REPLACEMENTS MAPPING ---
    replacements = {
        "A PROJECT PHASE I REPORT": "Makabulla Project Phase I Report",
        "HAZARD WASTE DETECTION USING AI AND IOT": meta["project_title"],
        "MVJ NEXTGEN EDUCLOUD:AI-BASED ACADEMIC RESOURCE MANAGER": meta["project_subtitle"],
        "BACHELOR OF ENGINEERING": "BACHELOR OF ENGINEERING", # Usually static, or add to JSON if needed
        "COMPUTER SCIENCE & ENGINEERING": meta["department_name"].replace("DEPARTMENT OF ", ""), # or just use meta['department_abbr']
        "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING": meta["department_name"],
        "MVJ COLLEGE OF ENGINEERING": meta["college_name"],
        "BANGALORE - 67": meta["city_code"],
        "ACADEMIC YEAR 2024 - 2025": f"ACADEMIC YEAR {meta['academic_year']}",
        
        # Guide Details
        "Mrs. Sumanjali": guides["name"],
        "Sumanjali": guides["name"],
        "Assistant Professor, Department of CSE": f"{guides['designation']}, {guides['department']}",
    }

    # --- Replace text in paragraphs (FRONT PAGE ONLY) ---
    # We assume the front page is the first section or explicitly the first page logic.
    # In python-docx, checking exactly "page 1" is hard, but usually front page = first section or first N paragraphs.
    # Best bet: Iterate paragraphs in the first section's range if possible, or just break early.
    # Simple approach: The Front Page in this template is likely Section 1.
    
    first_section = doc.sections[0]
    
    # Text replacements in body paragraphs of the first section
    # Note: doc.sections[0].parts is not directly iterable for paragraphs like this.
    # We must match the paragraphs that belong to the front page.
    # A safer approximate for "Front Page" is to replace only until we hit a page break or section break.
    
    for para in doc.paragraphs:
        # Check if we moved past the front page (naive check: if we see "CERTIFICATE" or similar)
        if "CERTIFICATE" in para.text.upper():
            break
            
        for old, new in replacements.items():
            replace_text_runs(para, old, new)

    # --- Replace text inside tables (FRONT PAGE ONLY) ---
    # The first table is the student list on the front page.
    if doc.tables:
        student_table = doc.tables[0] # Only touch the first table
        for row in student_table.rows:
            for cell in row.cells:
                for old, new in replacements.items():
                    if old in cell.text:
                        cell.text = cell.text.replace(old, new)
    
    # (Skip looping through all other tables)

    # --- Update student IDs and names ---
    if doc.tables:
        table = doc.tables[0]
        students = data["people"]["students"]
        for i, stud in enumerate(students):
            if i < len(table.rows):

                # Update USN with bold formatting
                cell_usn = table.cell(i, 0)
                cell_usn.text = "" # Clear existing content
                run_usn = cell_usn.paragraphs[0].add_run(stud["usn"])
                run_usn.bold = True
                
                # Update Name with bold formatting
                cell_name = table.cell(i, 1)
                cell_name.text = "" # Clear existing content
                run_name = cell_name.paragraphs[0].add_run(stud["name"])
                run_name.bold = True

    # --- HEADER (FRONT PAGE ONLY) ---
    # Usually Front Page has a separate header/footer or First Page Header.
    # We will only touch the first section's header.
    section = doc.sections[0] 
    header = section.header
    # Sometimes the header text is split or complex, but for this template it's just the title
    if header.paragraphs:
        replace_text_runs(header.paragraphs[0], "HAZARD WASTE DETECTION USING AI AND IOT", meta["project_title"])
    else:
        header.add_paragraph(meta["project_title"])

    # --- FOOTER (FRONT PAGE ONLY) ---
    # Handle various dash characters: regular hyphen (-), en dash (–), em dash (—)
    # Again, only touching Section 0
    footer = section.footer
    for para in footer.paragraphs:
        if not para.runs:
            continue
            
        # Get full text to check what's actually there
        full_text = "".join(run.text for run in para.runs)
        
        # Replace year patterns while preserving tabs and structure
        year_patterns = [
            (r"2024\s*[-–—]\s*2025", meta["academic_year"]),  # Matches "2024 - 2025", "2024 – 2025", etc.
            (r"ACADEMIC YEAR\s+2024\s*[-–—]\s*2025", f"ACADEMIC YEAR {meta['academic_year']}"),
        ]
        
        text_replaced = False
        for pattern, replacement in year_patterns:
            if re.search(pattern, full_text):
                # Find which run contains the year text
                for run in para.runs:
                    if re.search(pattern, run.text):
                        # Replace only the year part, preserve tabs and other text
                        run.text = re.sub(pattern, replacement, run.text)
                        text_replaced = True
                        break
                if text_replaced:
                    break
        
        # Replace department names if present
        dept_replacements = {
            "COMPUTER SCIENCE & ENGINEERING": meta["department_name"].replace("DEPARTMENT OF ", ""),
            "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING": meta["department_name"],
            "CSE": meta["department_abbr"],
            "Computer Science": meta["department_name"],
        }
        
        for old, new in dept_replacements.items():
            if old in full_text:
                # Replace in the run that contains the department text
                for run in para.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
                        break

    # Save updated doc
    doc.save(output_path)
    print(f"Dummy-fixed file saved as: {output_path}")

# --- RUN ---

if __name__ == "__main__":
    # Calculate input file path dynamically
    current_script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(os.path.dirname(os.path.dirname(current_script_dir)))
    input_file = os.path.join(project_root, "backend", "templates", "template_indivisual", "front_page.docx")
    
    if os.path.exists(input_file):
        update_project_doc_with_dummies(input_file)
    else:
        print(f"Error: Template file not found at {input_file}")