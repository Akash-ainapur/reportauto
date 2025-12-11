"""
Requirements Chapter Generator - Chapter 6

Generates Hardware and Software Requirements chapter.
Format: CHAPTER 6 -> HARDWARE AND SOFTWARE REQUIREMENTS -> 6.1 Software, 6.2 Hardware
"""

from docx import Document
from components.ai_content_base import AIContentBase


class RequirementsGenerator(AIContentBase):
    """
    Generates Chapter 6 - Hardware and Software Requirements.
    Lists technical specifications for project implementation.
    """
    NEEDS_BORDER = False

    def generate(self) -> Document:
        """Generate the Requirements chapter."""
        # Create fresh document
        doc = Document()
        
        # Add "CHAPTER 6"
        self.add_chapter_number(doc, 6)
        
        # Add "HARDWARE AND SOFTWARE REQUIREMENTS" as Heading 1
        self.add_heading(doc, "HARDWARE AND SOFTWARE REQUIREMENTS", level=1)
        
        # Add spacing
        doc.add_paragraph()
        
        # Generate AI content
        content = self.generate_ai_content("requirements")
        
        # Add content
        self.add_body_content(doc, content, justified=True, font_size=12.5)
        
        return doc
