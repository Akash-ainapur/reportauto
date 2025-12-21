from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .base import BaseSectionGenerator
import re

class CertificateGenerator(BaseSectionGenerator):
    NEEDS_BORDER = True

    def generate(self) -> Document:
        input_path = self.context.get_template_path("certificate.docx")
        doc = Document(input_path)
        self._process_certificate(doc)
        return doc

    def _process_certificate(self, doc):
        meta = self.meta
        cert = self.data["certificate_info"]

        # Styled Replacements
        styled_replacements = [
            ("MVJ COLLEGE OF ENGINEERING", meta.get("college_name", "MVJ COLLEGE OF ENGINEERING"), False, True, 18, None),
            ("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", meta.get("department_name", "DEPARTMENT OF COMPUTER SCIENCE").upper(), False, True, 14, None),
            ("Computer science and Engineering", meta.get("course_name", "Electronics and Communication Engineering"), False, False, None, None),
            (r"Major", meta.get("project_type", "Major"), True, False, None, None),
            ("2024 - 2025", meta.get("academic_year", "2024-2025"), False, False, None, None),
            ("Near ITPB, Whitefield, Bangalore – 560067", meta.get("college_address", "Near ITPB, Whitefield, Bangalore – 560067"), False, False, None, None),
            (r"[\"“]?\s*MVJ\s*NEXTGEN\s*EDUCLOUD[:\s-]+AI[:\s-]+BASED\s*ACADEMIC\s*RESOURCE\s*MANAGER\s*[\"”]?", 
             f'"{meta.get("project_title", "PROJECT TITLE")}"', True, True, None, RGBColor(255, 0, 0))
        ]

        for p in doc.paragraphs:
            for search, replace, is_re, bold, size, color in styled_replacements:
                 self._replace_chunk_with_format(p, search, replace, is_re, bold, size, color)

        self._rebuild_signature_block(doc, cert)

    def _rebuild_signature_block(self, doc, cert):
        # (Copied from original logic, simplified for class)
        # Try finding in table
        signature_row = None
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if "Signature of Guide" in cell.text:
                        signature_row = row
                        break
                if signature_row: break
            if signature_row: break

        def clean_text(text):
            if not text: return ""
            dept_short = self.meta.get("department_short", "Department of Universal Logic")
            text = text.replace("Department of CSE", dept_short)
            text = text.replace("Department of Universal Logic", dept_short)
            return text

        if signature_row:
            items = [
                ("Signature of Guide", clean_text(cert.get("guide_name", "[GUIDE NAME]")), clean_text(cert.get("guide_designation", "[DESIGNATION]"))),
                ("Signature of HOD",   clean_text(cert.get("hod_name", "[HOD NAME]")),   clean_text(cert.get("hod_designation", "[DESIGNATION]"))),
                ("Signature of Dean",  clean_text(cert.get("dean_name", "[DEAN NAME]")),  clean_text(cert.get("dean_designation", "[DESIGNATION]")))
            ]
            for col_idx, (header, name, designation) in enumerate(items):
                if col_idx >= len(signature_row.cells): break
                cell = signature_row.cells[col_idx]
                cell.text = ""
                # Header
                p = cell.paragraphs[0]
                r = p.add_run(header); r.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(2)
                # Name
                p = cell.add_paragraph()
                r = p.add_run(name); r.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(0)
                # Designation
                for line in designation.split('\n'):
                    if not line.strip(): continue
                    p = cell.add_paragraph()
                    r = p.add_run(line.strip()); r.bold = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            return

        # Fallback would be here but for now skipping as table exists in template

    def _replace_chunk_with_format(self, paragraph, old_text, new_text, is_regex=False, bold=False, size=None, color=None):
        # (Same logic as in original file, keeping it internal)
        if not paragraph.runs: return
        
        # 1. Capture
        run_data = []
        for r in paragraph.runs:
            run_data.append({
                "text": r.text, "bold": r.bold, "italic": r.italic, "underline": r.underline,
                "font_name": r.font.name, "font_size": r.font.size, "font_color": r.font.color.rgb,
            })
        full_text = "".join(d["text"] for d in run_data)
        
        # 2. Find
        match_start = -1
        match_end = -1
        if is_regex:
            match = re.search(old_text, full_text, flags=re.IGNORECASE)
            if match:
                match_start = match.start(); match_end = match.end()
        else:
            match_start = full_text.find(old_text)
            if match_start != -1: match_end = match_start + len(old_text)
        
        if match_start == -1: return

        # 3. Clear
        for r in paragraph.runs: r.text = ""

        # 4. Rebuild
        current_pos = 0
        for r_item in run_data:
            r_text = r_item["text"]
            r_len = len(r_text)
            r_end = current_pos + r_len
            
            def add_run(text):
                nr = paragraph.add_run(text)
                nr.bold = r_item["bold"]
                nr.italic = r_item["italic"]
                nr.underline = r_item["underline"]
                if r_item["font_name"]: nr.font.name = r_item["font_name"]
                if r_item["font_size"]: nr.font.size = r_item["font_size"]
                if r_item["font_color"]: nr.font.color.rgb = r_item["font_color"]
                return nr

            if r_end <= match_start or current_pos >= match_end:
                 add_run(r_text)
            else:
                if current_pos < match_start:
                    add_run(r_text[:match_start-current_pos])
                
                if current_pos <= match_start < r_end:
                    tr = paragraph.add_run(new_text)
                    tr.bold = True if bold else r_item["bold"]
                    if size: tr.font.size = Pt(size)
                    if color: tr.font.color.rgb = color
                
                if r_end > match_end:
                    add_run(r_text[match_end-current_pos:])
            current_pos += r_len
