from docx import Document
from docx.shared import Pt
from .base import BaseSectionGenerator
import re
from docx.enum.text import WD_ALIGN_PARAGRAPH

class AcknowledgementGenerator(BaseSectionGenerator):
    NEEDS_BORDER = True

    def generate(self) -> Document:
        input_path = self.context.get_template_path("Aknowledgement.docx")
        doc = Document(input_path)
        
        # Ensure ACKNOWLEDGEMENT is a Heading 1 for TOC indexing
        # Find the title paragraph and set style
        found_title = False
        for p in doc.paragraphs:
            if "ACKNOWLEDGEMENT" in p.text.upper():
                p.text = "ACKNOWLEDGEMENT"
                p.style = "Heading 1"
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                found_title = True
                break
        
        # If not found (unlikely), add it at the top
        if not found_title:
            new_p = doc.paragraphs[0].insert_paragraph_before("ACKNOWLEDGEMENT", style="Heading 1")
            new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        self._process_acknowledgement(doc)
        return doc

    def _process_acknowledgement(self, doc):
        approvers = self.data["people"].get("approvers", {})
        
        # Helper to safely get approver data
        def get_approver(role, field, default):
            return approvers.get(role, {}).get(field, default)

        # Safe Guide Access
        guides = self.data["people"].get("guide", [])
        if guides and len(guides) > 0:
            guide = guides[0]
        else:
            guide = {"name": "[GUIDE NAME]", "designation": "Assistant Professor", "department": "[DEPT]"}

        replacements = [
            (r"Dr\.\s*Ajayan\s*K\s*R", get_approver("principal", "name", "Dr. Ajayan K R"), True, True, None, None),
            (r"Principal\s+of\s+MVJCE", "Principal of MVJCE", True, True, None, None),
            (r"Dr\.\s*I\s*Hameem\s*Shanavas", get_approver("dean", "name", "Dr. I Hameem Shanavas"), True, True, None, None),
            (r"Dean,\s*School\s+of\s+CSE,\s*MVJCE", "Dean, School of CSE, MVJCE", True, True, None, None),
            (r"Mr\.\s*Kumar\s*R", get_approver("COE", "name", "Mr. Kumar R"), True, True, None, None),
            (r"Controller\s+Of\s+Examinations", "Controller Of Examinations", True, True, None, None),
            (r"Dr\.\s*Kiran\s*Babu\s*T\s*S", get_approver("hod", "name", "Dr. Kiran Babu T S"), True, True, None, None),
            (r"Mrs\.\s*Sumanjali\s*M", guide.get("name", "[GUIDE NAME]"), True, True, None, None),
            (r"Assistant\s+Professor|HEAD\s+WIZARD", guide.get("designation", "Assistant Professor"), True, True, None, None),
            (r"CSE\s+Department|ECE\s+Department", guide.get("department", "CSE Department"), True, True, None, None),
        ]

        self._process_paragraphs(doc.paragraphs, replacements)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                     self._process_paragraphs(cell.paragraphs, replacements)

    def _process_paragraphs(self, paragraphs, replacements):
        for p in paragraphs:
            for search, replace, is_re, bold, size, color in replacements:
                self._replace_chunk_with_format(p, search, replace, is_re, bold, size, color)

    def _replace_chunk_with_format(self, paragraph, old_text, new_text, is_regex=False, bold=False, size=None, color=None):
        # Using exact same logic as Certificate/Declaration
        if not paragraph.runs: return
        run_data = []
        for r in paragraph.runs:
            run_data.append({
                "text": r.text, "bold": r.bold, "italic": r.italic, "underline": r.underline,
                "font_name": r.font.name, "font_size": r.font.size, "font_color": r.font.color.rgb,
            })
        full_text = "".join(d["text"] for d in run_data)
        match_start = -1
        match_end = -1
        if is_regex:
            match = re.search(old_text, full_text, flags=re.IGNORECASE)
            if match: match_start = match.start(); match_end = match.end()
        else:
            match_start = full_text.find(old_text)
            if match_start != -1: match_end = match_start + len(old_text)
        if match_start == -1: return

        for r in paragraph.runs: r.text = ""
        current_pos = 0
        for r_item in run_data:
            r_text = r_item["text"]
            r_len = len(r_text)
            r_end = current_pos + r_len
            
            def add_run(text):
                nr = paragraph.add_run(text)
                nr.bold = r_item["bold"]
                nr.italic = r_item["italic"]
                nr.underline = r_item["underline"]
                if r_item["font_name"]: nr.font.name = r_item["font_name"]
                if r_item["font_size"]: nr.font.size = r_item["font_size"]
                if r_item["font_color"]: nr.font.color.rgb = r_item["font_color"]
                return nr

            if r_end <= match_start or current_pos >= match_end:
                 add_run(r_text)
            else:
                if current_pos < match_start: add_run(r_text[:match_start-current_pos])
                if current_pos <= match_start < r_end:
                     tr = paragraph.add_run(new_text)
                     tr.bold = True if bold else r_item["bold"]
                     if size: tr.font.size = Pt(size)
                     if color: tr.font.color.rgb = color
                if r_end > match_end: add_run(r_text[match_end-current_pos:])
            current_pos += r_len
