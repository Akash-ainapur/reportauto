"""
References Section Generator

Generates IEEE-formatted references based on literature survey.
Format: REFERENCES (Heading 1) -> Numbered citation list
"""

from docx import Document
from components.ai_content_base import AIContentBase


class ReferencesGenerator(AIContentBase):
    """
    Generates References section.
    IEEE-formatted citations from literature survey.
    """
    NEEDS_BORDER = False

    def generate(self) -> Document:
        """Generate the References section."""
        # Create fresh document
        doc = Document()
        
        # Add "REFERENCES" as Heading 1
        self.add_heading(doc, "REFERENCES", level=1)
        
        # Add spacing
        doc.add_paragraph()
        
        # Generate AI content
        content = self.generate_ai_content("references")
        
        # Add content
        self.add_body_content(doc, content, justified=True, font_size=12.5)
        
        return doc
