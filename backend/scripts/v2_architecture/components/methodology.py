"""
Methodology Chapter Generator - Chapter 5

Generates the Methodology chapter with technical details.
Format: CHAPTER 5 -> METHODOLOGY (Heading 1) -> Technical approach with subsections
"""

from docx import Document
from components.ai_content_base import AIContentBase


class MethodologyGenerator(AIContentBase):
    """
    Generates Chapter 5 - Methodology.
    Technical implementation approach with system design details.
    """
    NEEDS_BORDER = False

    def generate(self) -> Document:
        """Generate the Methodology chapter."""
        # Create fresh document
        doc = Document()
        
        # Add "CHAPTER 5"
        self.add_chapter_number(doc, 5)
        
        # Add "METHODOLOGY" as Heading 1
        self.add_heading(doc, "METHODOLOGY", level=1)
        
        # Add spacing
        doc.add_paragraph()
        
        # Generate AI content
        content = self.generate_ai_content("methodology")
        
        # Add content
        self.add_body_content(doc, content, justified=True, font_size=12.5)
        
        return doc
