"""
Introduction Chapter Generator - Chapter 1

Generates the Introduction chapter with AI-powered content.
Format: CHAPTER 1 (18pt bold) -> INTRODUCTION (Heading 1) -> Content (justified)
"""

from docx import Document
from components.ai_content_base import AIContentBase


class IntroductionGenerator(AIContentBase):
    """
    Generates Chapter 1 - Introduction with AI-powered content.
    Follows the document structure: Chapter number -> Chapter title -> Content
    """
    NEEDS_BORDER = False

    def generate(self) -> Document:
        """Generate the Introduction chapter with AI content."""
        # Create fresh document
        doc = Document()
        
        # Add "CHAPTER 1" (18pt, bold, left-aligned)
        self.add_chapter_number(doc, 1)
        
        # Add "INTRODUCTION" as Heading 1 (for TOC indexing)
        self.add_heading(doc, "INTRODUCTION", level=1)
        
        # Add spacing
        doc.add_paragraph()
        
        # Generate AI content
        content = self.generate_ai_content("introduction")
        
        # Add content with justified alignment
        self.add_body_content(doc, content, justified=True, font_size=13)
        
        return doc
