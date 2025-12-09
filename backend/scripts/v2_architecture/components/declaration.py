from docx import Document
from docx.shared import Pt, RGBColor
from .base import BaseSectionGenerator
import re

class DeclarationGenerator(BaseSectionGenerator):
    NEEDS_BORDER = True

    def generate(self) -> Document:
        input_path = self.context.get_template_path("declaration.docx")
        doc = Document(input_path)
        self._process_declaration(doc)
        return doc

    def _process_declaration(self, doc):
        meta = self.meta
        students = self.data["people"]["students"]
        
        def get_stud_name(idx): return students[idx]["name"] if idx < len(students) else f"Student {idx+1}"
        def get_stud_usn(idx): return students[idx]["usn"] if idx < len(students) else f"USN {idx+1}"

        replacements = [
            ("MVJ COLLEGE OF ENGINEERING", meta["college_name"], False, True, 18, None),
            ("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", meta["department_name"], False, True, 14, None),
            ("Computer science and Engineering", meta["department_short"], False, True, 12, None),
            (r"Sixth", meta["semester"], True, False, None, None),
            (r"Major", meta["project_type"], True, False, None, None),
            ("ARUNPRAKASH PATTAR", get_stud_name(0), False, False, None, None),
            ("HEMANTH S SAVANTH",  get_stud_name(1), False, False, None, None),
            ("CHINMAHESH C",       get_stud_name(2), False, False, None, None),
            ("BHAVYASHREE R",      get_stud_name(3), False, False, None, None),
            ("1MJ22CS217", get_stud_usn(0), False, False, None, None),
            ("1MJ23CS411", get_stud_usn(1), False, False, None, None),
            ("1MJ23CS406", get_stud_usn(2), False, False, None, None),
            ("1MJ22CS029", get_stud_usn(3), False, False, None, None),
            ("2024-2025", meta["academic_year"], False, False, None, None),
            ("2024 - 2025", meta["academic_year"], False, False, None, None),
             ("Computer science and Engineering", meta.get("department_short", "Dept of Universal Logic"), False, False, None, None),
            (r"[\"“]?\s*MVJ\s*NEXTGEN\s*EDUCLOUD[:\s-]+AI[:\s-]+BASED\s*ACADEMIC\s*RESOURCE\s*MANAGER\s*[\"”]?", 
             [
                 {"text": "“", "bold": True, "color": None},
                 {"text": meta["project_title"], "bold": True, "color": RGBColor(255, 0, 0)},
                 {"text": "”", "bold": True, "color": None},
             ], 
             True, True, None, RGBColor(255, 0, 0)),
        ]

        # Process Body
        self._process_paragraphs(doc.paragraphs, replacements)
        # Process Tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                     self._process_paragraphs(cell.paragraphs, replacements)

    def _process_paragraphs(self, paragraphs, replacements):
        for p in paragraphs:
            for search, replace, is_re, bold, size, color in replacements:
                self._replace_chunk_with_format(p, search, replace, is_re, bold, size, color)

    def _replace_chunk_with_format(self, paragraph, old_text, new_text, is_regex=False, bold=False, size=None, color=None):
        # Using exact same logic as Certificate
        if not paragraph.runs: return
        run_data = []
        for r in paragraph.runs:
            run_data.append({
                "text": r.text, "bold": r.bold, "italic": r.italic, "underline": r.underline,
                "font_name": r.font.name, "font_size": r.font.size, "font_color": r.font.color.rgb,
            })
        full_text = "".join(d["text"] for d in run_data)
        match_start = -1
        match_end = -1
        if is_regex:
            match = re.search(old_text, full_text, flags=re.IGNORECASE)
            if match: match_start = match.start(); match_end = match.end()
        else:
            match_start = full_text.find(old_text)
            if match_start != -1: match_end = match_start + len(old_text)
        if match_start == -1: return

        for r in paragraph.runs: r.text = ""
        current_pos = 0
        for r_item in run_data:
            r_text = r_item["text"]
            r_len = len(r_text)
            r_end = current_pos + r_len
            
            def add_run(text):
                nr = paragraph.add_run(text)
                nr.bold = r_item["bold"]
                nr.italic = r_item["italic"]
                nr.underline = r_item["underline"]
                if r_item["font_name"]: nr.font.name = r_item["font_name"]
                if r_item["font_size"]: nr.font.size = r_item["font_size"]
                if r_item["font_color"]: nr.font.color.rgb = r_item["font_color"]
                return nr

            if r_end <= match_start or current_pos >= match_end:
                 add_run(r_text)
            else:
                if current_pos < match_start: add_run(r_text[:match_start-current_pos])
                if current_pos <= match_start < r_end:
                    if isinstance(new_text, list):
                        for seg in new_text:
                            tr = paragraph.add_run(seg.get("text", ""))
                            tr.bold = True if seg.get("bold", bold) else r_item["bold"]
                            if size: tr.font.size = Pt(size)
                            if seg.get("color", color): tr.font.color.rgb = seg.get("color", color)
                    else:
                        tr = paragraph.add_run(new_text)
                        tr.bold = True if bold else r_item["bold"]
                        if size: tr.font.size = Pt(size)
                        if color: tr.font.color.rgb = color
                if r_end > match_end: add_run(r_text[match_end-current_pos:])
            current_pos += r_len
