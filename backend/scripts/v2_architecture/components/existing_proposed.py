"""
Existing and Proposed System Chapter Generator - Chapter 4

Generates Chapter 4 with comparative analysis.
Format: CHAPTER 4 -> Heading 1 -> 4.1 Existing, 4.1.1 Disadvantages, 4.2 Proposed, 4.2.1 Advantages
"""

from docx import Document
from components.ai_content_base import AIContentBase


class ExistingProposedGenerator(AIContentBase):
    """
    Generates Chapter 4 - Existing and Proposed System.
    Includes detailed comparison with numbered subsections.
    """
    NEEDS_BORDER = False

    def generate(self) -> Document:
        """Generate the Existing and Proposed System chapter."""
        # Create fresh document
        doc = Document()
        
        # Add "CHAPTER 4: EXISTING AND PROPOSED SYSTEM" as Heading 1
        self.add_heading(doc, "CHAPTER 4: EXISTING AND PROPOSED SYSTEM", level=1)
        
        # Add spacing
        doc.add_paragraph()
        
        # Generate AI content
        content = self.generate_ai_content("existing_proposed")
        
        # Add content
        self.add_body_content(doc, content, justified=True, font_size=12.5)
        
        return doc
