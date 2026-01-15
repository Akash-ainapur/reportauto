from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from components.ai_content_base import AIContentBase


class AbstractGenerator(AIContentBase):
    """
    Generates the Abstract page with AI-powered content.
    Format: ABSTRACT heading (18pt bold, centered) followed by justified content.
    """
    NEEDS_BORDER = True

    def generate(self) -> Document:
        """Generate the abstract page with AI content."""
        # Create fresh document
        doc = Document()
        
        # Add ABSTRACT title (Heading 1 for TOC, centered)
        self.add_heading(doc, "ABSTRACT", level=1, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        
        # Add spacing
        doc.add_paragraph()
        
        # Generate AI content
        content = self.generate_ai_content("abstract")
        
        # Add content with justified alignment
        self.add_body_content(doc, content, justified=True, font_size=12)
        
        return doc
