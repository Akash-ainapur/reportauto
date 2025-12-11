"""
Literature Survey Chapter Generator - Chapter 2

Generates the Literature Survey chapter with AI-powered content.
Format: CHAPTER 2 (18pt bold) -> LITERATURE SURVEY (Heading 1) -> Numbered entries (2.1, 2.2, etc.)
"""

from docx import Document
from components.ai_content_base import AIContentBase


class LiteratureSurveyGenerator(AIContentBase):
    """
    Generates Chapter 2 - Literature Survey with AI-powered content.
    Includes 10-14 numbered subsections with research paper summaries.
    """
    NEEDS_BORDER = False

    def generate(self) -> Document:
        """Generate the Literature Survey chapter with AI content."""
        # Create fresh document
        doc = Document()
        
        # Add "CHAPTER 2" (18pt, bold, left-aligned)
        self.add_chapter_number(doc, 2)
        
        # Add "LITERATURE SURVEY" as Heading 1 (for TOC indexing)
        self.add_heading(doc, "LITERATURE SURVEY", level=1)
        
        # Add spacing
        doc.add_paragraph()
        
        # Generate AI content
        content = self.generate_ai_content("literature_survey")
        
        # Add content with justified alignment
        self.add_body_content(doc, content, justified=True, font_size=12.5)
        
        return doc
