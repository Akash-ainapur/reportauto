from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .base import BaseSectionGenerator
import re

class FrontPageGenerator(BaseSectionGenerator):
    NEEDS_BORDER = False

    def generate(self) -> Document:
        input_path = self.context.get_template_path("front_page.docx")
        doc = Document(input_path)
        
        self._process_front_page(doc)
        return doc

    def _process_front_page(self, doc):
        meta = self.meta
        
        # Safe Guide Extraction
        guides = self.data["people"].get("guide", [])
        if guides and len(guides) > 0:
            guide = guides[0]
        else:
            guide = {
                "name": "[GUIDE NAME]",
                "designation": "[DESIGNATION]",
                "department": "[DEPARTMENT]"
            }
            
        styles = self.data.get("styles", {})
        
        # Replacements Mapping
        replacements = {
            "A PROJECT PHASE I REPORT": "Makabulla Project Phase I Report",
            "HAZARD WASTE DETECTION USING AI AND IOT": meta.get("project_title", "PROJECT TITLE"),
            "MVJ NEXTGEN EDUCLOUD:AI-BASED ACADEMIC RESOURCE MANAGER": meta.get("project_subtitle", ""),
            "BACHELOR OF ENGINEERING": "BACHELOR OF ENGINEERING",
            "COMPUTER SCIENCE & ENGINEERING": meta.get("department_name", "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING").replace("DEPARTMENT OF ", ""),
            "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING": meta.get("department_name", "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING"),
            "MVJ COLLEGE OF ENGINEERING": meta.get("college_name", "MVJ COLLEGE OF ENGINEERING"),
            "BANGALORE - 67": meta.get("city_code", "BANGALORE - 67"),
            "ACADEMIC YEAR 2024 - 2025": f"ACADEMIC YEAR {meta.get('academic_year', '2024-25')}",
            "Assistant Professor, Department of CSE": f"{guide['designation'].title()}, {guide['department'].title()}",
            "Under the Guidance of": "Under the Guidance of",
        }

        guide_name_placeholders = ["Mrs. Sumanjali", "Sumanjali"]
        guide_name_new = guide["name"]
        
        # Styles
        placeholder_styles = {
            "HAZARD WASTE DETECTION USING AI AND IOT": styles.get("project_title"),
            "MVJ NEXTGEN EDUCLOUD:AI-BASED ACADEMIC RESOURCE MANAGER": styles.get("project_subtitle"),
            "Under the Guidance of": styles.get("guide_label"),
        }
        guide_style = styles.get("guide_name")

        # 1. Process Paragraphs (Body)
        for para in doc.paragraphs:
            if "CERTIFICATE" in para.text.upper():
                break
            for old, new in replacements.items():
                font_sz = placeholder_styles.get(old)
                self._replace_text_runs(para, old, new, font_size=font_sz)
            for old_name in guide_name_placeholders:
                self._replace_with_bold(para, old_name, guide_name_new, font_size=guide_style)

        # 2. Process Tables (Student List)
        if doc.tables:
            student_table = doc.tables[0]
            # Text Replacements in table
            for row in student_table.rows:
                for cell in row.cells:
                    for old, new in replacements.items():
                        if old in cell.text:
                            cell.text = cell.text.replace(old, new)
            
            # Student Details
            students = self.data["people"]["students"]
            student_font_size = styles.get("student_details")
            for i, stud in enumerate(students):
                if i < len(student_table.rows):
                    self._update_student_row(student_table.rows[i], stud, student_font_size)

        # 3. Header/Footer
        section = doc.sections[0]
        if section.header.paragraphs:
            self._replace_text_runs(section.header.paragraphs[0], "HAZARD WASTE DETECTION USING AI AND IOT", meta["project_title"])
        
        self._process_footer(section.footer)

    def _process_footer(self, footer):
        meta = self.meta
        for para in footer.paragraphs:
            if not para.runs: continue
            
            full_text = "".join(r.text for r in para.runs)
            
            year_patterns = [
                (r"2024\s*[-–—]\s*2025", meta["academic_year"]),
                (r"ACADEMIC YEAR\s+2024\s*[-–—]\s*2025", f"ACADEMIC YEAR {meta['academic_year']}"),
            ]
            for pattern, replacement in year_patterns:
                if re.search(pattern, full_text):
                    for run in para.runs:
                        if re.search(pattern, run.text):
                            run.text = re.sub(pattern, replacement, run.text)
            
            dept_replacements = {
                "COMPUTER SCIENCE & ENGINEERING": meta.get("department_name", "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING").replace("DEPARTMENT OF ", ""),
                "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING": meta.get("department_name", "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING"),
                "CSE": meta.get("department_abbr", "CSE"),
                "Computer Science": meta.get("department_name", "Computer Science"),
            }
            for old, new in dept_replacements.items():
                if old in full_text:
                    for run in para.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)

    def _update_student_row(self, row, stud, font_size):
        # USN
        cell = row.cells[0]
        if cell.paragraphs:
            p = cell.paragraphs[0]
            for r in p.runs: r.text = ""
            run = p.add_run(stud["usn"])
            run.bold = True
            if font_size: run.font.size = Pt(font_size)
        # Name
        cell = row.cells[1]
        if cell.paragraphs:
            p = cell.paragraphs[0]
            for r in p.runs: r.text = ""
            run = p.add_run(stud["name"])
            run.bold = True
            if font_size: run.font.size = Pt(font_size)

    def _replace_text_runs(self, paragraph, old, new, font_size=None):
        full_text = "".join(run.text for run in paragraph.runs)
        if old in full_text:
            new_text = full_text.replace(old, new)
            for run in paragraph.runs: run.text = ""
            run = paragraph.runs[0]
            run.text = new_text
            if font_size: run.font.size = Pt(font_size)

    def _replace_with_bold(self, paragraph, old, new, font_size=None):
        full_text = "".join(run.text for run in paragraph.runs)
        if old in full_text:
            parts = full_text.split(old)
            for run in paragraph.runs: run.text = ""
            for i, part in enumerate(parts):
                paragraph.add_run(part)
                if i < len(parts) - 1:
                    run = paragraph.add_run(new)
                    run.bold = True
                    if font_size: run.font.size = Pt(font_size)
