"""
Base class for AI-powered content generators.

Provides common functionality for components that generate content using AI,
including proper formatting, heading styles, and content generation.
"""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys
import os

# Add parent directory to path for imports
current_dir = os.path.dirname(os.path.abspath(__file__))
parent_dir = os.path.dirname(current_dir)
if parent_dir not in sys.path:
    sys.path.insert(0, parent_dir)

from components.base import BaseSectionGenerator
from services.gemini_service import GeminiService
from services.prompt_templates import PromptTemplates


class AIContentBase(BaseSectionGenerator):
    """
    Base class for AI-generated content sections.
    Provides helper methods for consistent formatting and AI content generation.
    """
    
    def __init__(self, context):
        super().__init__(context)
        self.ai_service = GeminiService()
        self.ai_config = self.data.get("ai_content_config", {})
    
    def add_chapter_number(self, doc: Document, chapter_num: int) -> None:
        """
        Add chapter number in the correct format.
        Format: "CHAPTER X" - Normal style, 18pt, Bold, Left-aligned
        
        Args:
            doc: Document object
            chapter_num: Chapter number (e.g., 1, 2, 3)
        """
        p = doc.add_paragraph()
        run = p.add_run(f"CHAPTER {chapter_num}")
        run.bold = True
        run.font.size = Pt(18)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    def add_header_footer(self, doc: Document) -> None:
        """
        Add header and footer to the document.
        
        Header: Project title (CAPS, left-aligned, 12pt)
        Footer: Department (left) | Page number (center) | Year (right) - all 12pt
        """
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.shared import OxmlElement
        from docx.oxml.ns import qn
        
        # Get data from JSON
        global_meta = self.data.get("global_metadata", {})
        project_title = global_meta.get("project_title", "").upper()
        department_name = global_meta.get("department_name", "")
        academic_year = global_meta.get("academic_year", "")
        
        # Access the default section
        section = doc.sections[0]
        
        # === HEADER ===
        header = section.header
        header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        header_para.text = project_title
        header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # Format header text
        for run in header_para.runs:
            run.font.size = Pt(12)
        
        # === FOOTER ===
        footer = section.footer
        footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        footer_para.clear()
        
        # Left: Department name
        left_run = footer_para.add_run(department_name)
        left_run.font.size = Pt(12)
        
        # Add tab for center alignment
        footer_para.add_run("\t")
        
        # Center: Page number
        center_run = footer_para.add_run()
        center_run.font.size = Pt(12)
        
        # Add page number field
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        center_run._element.append(fldChar1)
        
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        center_run._element.append(instrText)
        
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        center_run._element.append(fldChar2)
        
        # Add tab for right alignment
        footer_para.add_run("\t")
        
        # Right: Year
        right_run = footer_para.add_run(academic_year)
        right_run.font.size = Pt(12)
        
        # Set tab stops for footer (left, center, right)
        from docx.shared import Inches
        tab_stops = footer_para.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(3.25), WD_ALIGN_PARAGRAPH.CENTER)
        tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)
    
    def add_heading(self, doc: Document, text: str, level: int = 1) -> None:
        """
        Add a heading with the correct style for indexing.
        
        Args:
            doc: Document object
            text: Heading text
            level: Heading level (1 for H1, 2 for H2, 3 for H3)
        """
        heading_style = f"Heading {level}"
        heading = doc.add_heading(text, level=level)
        # Ensure uppercase for main headings (level 1)
        if level == 1:
            heading.text = text.upper()
        
        # Apply 1.5 line spacing to all headings
        heading.paragraph_format.line_spacing = 1.5
    
    def add_body_content(self, doc: Document, text: str, justified: bool = True, font_size: float = 12.5, line_spacing: float = 1.5) -> None:
        """
        Add body content with proper formatting, parsing markdown syntax.
        
        Supports:
        - ### Heading (converted to Heading 3)
        - **bold text** (converted to bold)
        - *italic text* (converted to italic)
        
        Args:
            doc: Document object
            text: Content text (may contain markdown)
            justified: Whether to justify the text (default: True)
            font_size: Font size in points (default: 12.5, can be decimal)
            line_spacing: Line spacing multiplier (default: 1.5 for academic formatting)
        """
        import re
        from docx.shared import Pt
        
        # Split text into lines for processing
        lines = text.split('\n')
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            # Check if line is a markdown heading
            heading_match = re.match(r'^(#{1,3})\s+(.+)$', line)
            if heading_match:
                level_symbols = heading_match.group(1)
                heading_text = heading_match.group(2)
                
                # Determine heading level (### = H3, ## = H2, # = H1)
                # But we want to use H2 and H3 for subsections within chapters
                heading_level = len(level_symbols)
                if heading_level == 1:
                    heading_level = 2  # Convert # to H2
                elif heading_level == 2:
                    heading_level = 2  # Convert ## to H2
                else:
                    heading_level = 3  # Convert ### to H3
                
                # Add as proper heading
                self.add_heading(doc, heading_text, level=heading_level)
                continue
            
            # Process line with inline markdown (bold, italic)
            p = doc.add_paragraph()
            if justified:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            # Set line spacing
            p.paragraph_format.line_spacing = line_spacing
            
            # Parse inline formatting: **bold** and *italic*
            # We'll use a simple regex-based approach
            current_pos = 0
            
            # Pattern to match **bold** or *italic* or regular text
            pattern = r'(\*\*(.+?)\*\*|\*(.+?)\*|[^*]+|\*)'
            
            for match in re.finditer(pattern, line):
                matched_text = match.group(0)
                
                if matched_text.startswith('**') and matched_text.endswith('**'):
                    # Bold text
                    bold_text = match.group(2)
                    run = p.add_run(bold_text)
                    run.bold = True
                    run.font.size = Pt(font_size)
                elif matched_text.startswith('*') and matched_text.endswith('*') and not matched_text.startswith('**'):
                    # Italic text
                    italic_text = match.group(3)
                    run = p.add_run(italic_text)
                    run.italic = True
                    run.font.size = Pt(font_size)
                elif matched_text and matched_text != '*':
                    # Regular text
                    run = p.add_run(matched_text)
                    run.font.size = Pt(font_size)
    
    def generate_ai_content(self, section_name: str) -> str:
        """
        Generate AI content for a specific section.
        
        Args:
            section_name: Name of the section (e.g., 'abstract', 'introduction')
            
        Returns:
            Generated content text
        """
        # Check if AI generation is enabled globally
        if not self.ai_config.get("enabled", False):
            return self._get_manual_content(section_name)
        
        # Check if AI generation is enabled for this specific section
        section_config = self.ai_config.get("sections", {}).get(section_name, {})
        if not section_config.get("generate", False):
            return self._get_manual_content(section_name)
        
        # Get project context from JSON data
        global_meta = self.data.get("global_metadata", {})
        
        # Prepare context for prompt template
        context = {
            "project_title": global_meta.get("project_title", "Untitled Project"),
            "project_type": global_meta.get("project_type", "engineering"),
            "domain": global_meta.get("domain", "technology"),
            "word_limit": section_config.get("word_limit", PromptTemplates.get_word_limit(section_name))
        }
        
        # Check for custom prompt
        custom_prompt = section_config.get("custom_prompt")
        if custom_prompt:
            # Use custom prompt directly
            system_instruction = PromptTemplates.SYSTEM_INSTRUCTIONS.get(section_name, "")
            prompt = custom_prompt.format(**context)
        else:
            # Use template from PromptTemplates
            system_instruction, prompt = PromptTemplates.get_prompt(section_name, context)
        
        # Get temperature from config
        temperature = self.ai_config.get("temperature", 0.7)
        
        # Generate content
        content = self.ai_service.generate_content(
            prompt=prompt,
            system_instruction=system_instruction,
            temperature=temperature
        )
        
        return content
    
    def _get_manual_content(self, section_name: str) -> str:
        """
        Get manually written content from JSON if available.
        
        Args:
            section_name: Name of the section
            
        Returns:
            Manual content or placeholder
        """
        report_content = self.data.get("report_content", {})
        section_data = report_content.get(section_name, {})
        
        # Check for 'text' key in section data
        if isinstance(section_data, dict):
            return section_data.get("text", f"[{section_name.title()} content placeholder]")
        
        return f"[{section_name.title()} content placeholder]"
