"""
Test script to verify header/footer in Introduction page
"""

import sys
from pathlib import Path

# Add parent directory to path
current_dir = Path(__file__).resolve().parent
project_root = current_dir.parent.parent.parent.parent
sys.path.insert(0, str(project_root / 'backend' / 'scripts' / 'v2_architecture'))

from docx import Document
from components.introduction import IntroductionGenerator
from builder import ReportContext

print("=" * 60)
print("Testing Header/Footer on Introduction Page")
print("=" * 60)

# Create context
context = ReportContext(str(project_root))

# Generate introduction
generator = IntroductionGenerator(context)
intro_doc = generator.generate()

# Check header
section = intro_doc.sections[0]
header = section.header
print(f"\nHeader paragraphs: {len(header.paragraphs)}")
if header.paragraphs:
    print(f"Header text: '{header.paragraphs[0].text}'")

# Check footer
footer = section.footer
print(f"\nFooter paragraphs: {len(footer.paragraphs)}")
if footer.paragraphs:
    print(f"Footer text: '{footer.paragraphs[0].text}'")

# Save standalone test
test_output = project_root / 'output' / 'test_introduction_only.docx'
intro_doc.save(str(test_output))

print(f"\n✅ Saved standalone test to: {test_output}")
print("\nOpen this file to verify headers/footers appear correctly.")
print("If they appear here but not in full_report_v2.docx, the issue")
print("is with document merging in the builder.")
