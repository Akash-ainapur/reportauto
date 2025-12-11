"""
Script to analyze Introduction template structure
"""
from docx import Document
import sys

template_path = r"d:\Model\REPORT1\AutomatedReportGenerator\backend\templates\template_indivisual\senior_project_template (10).docx"

try:
    doc = Document(template_path)
    
    print("=" * 70)
    print("ANALYZING INTRODUCTION TEMPLATE STRUCTURE")
    print("=" * 70)
    
    # Look for Introduction section
    found_intro = False
    intro_content = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        
        # Check if we found introduction
        if "INTRODUCTION" in text.upper() or "CHAPTER" in text.upper():
            found_intro = True
            print(f"\n✓ Found: {text}")
            print(f"  Style: {para.style.name}")
            continue
        
        # Collect content after introduction
        if found_intro and text:
            intro_content.append({
                'text': text[:100] + "..." if len(text) > 100 else text,
                'style': para.style.name
            })
            
            # Stop after collecting ~15 paragraphs
            if len(intro_content) > 15:
                break
    
    print("\n" + "=" * 70)
    print("INTRODUCTION CONTENT STRUCTURE:")
    print("=" * 70)
    
    for i, item in enumerate(intro_content, 1):
        print(f"\n{i}. Style: {item['style']}")
        print(f"   Text: {item['text']}")
    
    print("\n" + "=" * 70)
    print("HEADINGS IN DOCUMENT:")
    print("=" * 70)
    
    for para in doc.paragraphs:
        if para.style.name.startswith('Heading'):
            print(f"\n{para.style.name}: {para.text}")

except Exception as e:
    print(f"Error: {e}")
    import traceback
    traceback.print_exc()
