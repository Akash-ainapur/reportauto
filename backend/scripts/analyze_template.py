from docx import Document
import re

def detailed_scan(path):
    print(f"--- DETAILED SCAN: {path} ---")
    doc = Document(path)
    
    keywords = ["Principal", "Date", "Place", "External", "Certify", "Declare", "Hardware", "Software"]
    
    all_text = []
    
    # Collect all text first
    for para in doc.paragraphs:
        all_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text.append(cell.text)
                
    full_content = "\n".join(all_text)
    
    print("\n--- KEYWORD HITS ---")
    for kw in keywords:
        matches = [line for line in all_text if kw.lower() in line.lower()]
        if matches:
            print(f"FOUND '{kw}':")
            for m in matches[:3]: # Limit verification output
                print(f"  > {m.strip()[:100]}...")
        else:
            print(f"NOT FOUND: '{kw}'")

    print("\n--- TABLES STRUCTURE (First few rows) ---")
    for i, table in enumerate(doc.tables):
        print(f"Table {i}:")
        for r, row in enumerate(table.rows[:3]):
             print(f"  R{r}: {[c.text.strip()[:20] for c in row.cells]}")

if __name__ == "__main__":
    path = r"d:\Model\REPORT1\AutomatedReportGenerator\backend\templates\senior_project_template.docx"
    detailed_scan(path)
